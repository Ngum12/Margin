"""Preview, then confirm a register. Codes preserve identity; names never auto-merge."""

import csv
import io
import zipfile
from pathlib import Path
from typing import Any

from PIL import Image
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from app.ai.factory import ollama_client
from app.ai.ollama import resolve_model
from app.ai.prompts import load_prompt
from app.config import Settings
from app.db.models import Exam, JobType, RosterImport, Script, Student, User
from app.errors import ConflictError, NotFoundError, ValidationFailedError
from app.pdf.runner import render_pdf
from app.schemas.roster import RosterRow, RosterRows
from app.services import audit, jobs, students
from app.services.access import get_visible_class, get_visible_student

MAX_BYTES = 20 * 1024 * 1024
SUFFIXES = {".csv", ".txt", ".xlsx", ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".pdf"}


def class_students(db: Session, class_id: int) -> list[Student]:
    return list(
        db.scalars(
            select(Student)
            .where(Student.class_id == class_id, Student.is_active.is_(True))
            .order_by(Student.register_position, Student.id)
        )
    )


def exam_students(db: Session, exam: Exam, *, refresh: bool = False) -> list[Student]:
    if exam.roster_student_ids is None or (
        refresh and not db.scalar(select(Script.id).where(Script.exam_id == exam.id).limit(1))
    ):
        exam.roster_student_ids = [s.id for s in class_students(db, exam.class_id)]
        db.flush()
    result = {
        s.id: s for s in db.scalars(select(Student).where(Student.id.in_(exam.roster_student_ids)))
    }
    return [result[ident] for ident in exam.roster_student_ids if ident in result]


def visible(db: Session, actor: User, ident: int) -> RosterImport:
    item = db.get(RosterImport, ident)
    if not item:
        raise NotFoundError("Register import")
    get_visible_class(db, actor, item.class_id)
    return item


def parse_rows(rows: list[list[str]]) -> list[dict[str, Any]]:
    rows = [r for r in rows if any(v.strip() for v in r)]
    if not rows:
        raise ValueError("empty")
    header = [v.strip().lower() for v in rows[0]]
    names = {"name", "nom", "student", "élève", "student name", "nom de l'élève"}
    codes = {"code", "student code", "code élève"}
    name_col = next((i for i, v in enumerate(header) if v in names), None)
    code_col = next((i for i, v in enumerate(header) if v in codes), None)
    if name_col is not None:
        rows = rows[1:]
    else:
        name_col = 0
    parsed = []
    for row in rows:
        name = row[name_col] if len(row) > name_col else ""
        code = row[code_col] if code_col is not None and len(row) > code_col else ""
        parsed.append(RosterRow(name=name, code=code))
    return [r.model_dump() for r in RosterRows(rows=parsed).rows]


def spreadsheet(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    with zipfile.ZipFile(path) as zipped:
        entries = zipped.infolist()
        if len(entries) > 2000 or sum(i.file_size for i in entries) > 50 * 1024 * 1024:
            raise ValueError("archive_size")
        if any(
            "externallinks/" in i.filename.lower() or "vbaproject" in i.filename.lower()
            for i in entries
        ):
            raise ValueError("external_content")
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheet = book.worksheets[0]
        rows = []
        for cells in sheet.iter_rows(max_row=302, max_col=20):
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError("formula")
            rows.append([str(cell.value) if cell.value is not None else "" for cell in cells])
        return parse_rows(rows)
    finally:
        book.close()


def receive(
    db: Session, actor: User, class_id: int, content: bytes, suffix: str, settings: Settings
) -> RosterImport:
    get_visible_class(db, actor, class_id)
    suffix = suffix.lower()
    if suffix not in SUFFIXES or not content or len(content) > MAX_BYTES:
        raise ValidationFailedError("Unsupported or oversized register", code="roster.invalid")
    item = RosterImport(class_id=class_id, created_by_id=actor.id, file_path="", suffix=suffix)
    db.add(item)
    db.flush()
    item.file_path = f"roster-imports/{item.id}/original{suffix}"
    path = settings.data_dir / item.file_path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    jobs.enqueue(
        db, JobType.EXTRACT_ROSTER, {"roster_id": item.id}, f"roster:{item.id}", max_attempts=3
    )
    audit.record(
        db,
        actor=actor,
        entity_type="roster_import",
        entity_id=item.id,
        action="import_register",
        new_value={"class_id": class_id, "suffix": suffix},
    )
    return item


def extract(factory: sessionmaker[Session], settings: Settings, payload: dict[str, Any]) -> None:
    with factory() as db:
        item = db.get(RosterImport, int(payload["roster_id"]))
        if not item or item.status in {"ready", "confirmed"}:
            return
        path = settings.data_dir / item.file_path
        try:
            if item.suffix == ".xlsx":
                item.rows = spreadsheet(path)
            elif item.suffix in {".txt", ".csv"}:
                item.text = path.read_text(encoding="utf-8-sig")
                if len(item.text) > 100_000:
                    raise ValueError("text_size")
                delimiter = (
                    "\t" if "\t" in item.text else ";" if ";" in item.text.splitlines()[0] else ","
                )
                item.rows = parse_rows(
                    [[line] for line in item.text.splitlines()]
                    if item.suffix == ".txt" and "\t" not in item.text
                    else list(csv.reader(io.StringIO(item.text), delimiter=delimiter))
                )
            else:
                if settings.ai_backend == "mock":
                    item.status, item.error = "manual", "scheme.local_model_required"
                    db.commit()
                    return
                rendered = render_pdf(path, path.parent / "pages", settings)
                client = ollama_client(settings)
                model = resolve_model(client, settings.reader_model)
                prompt = load_prompt("register_read")
                item.model_name, item.model_version, item.prompt_version = (
                    model.name,
                    model.version,
                    prompt.version,
                )
                all_rows = []
                for page in rendered.pages:
                    cache = path.parent / f"read-{page.index}.json"
                    if cache.exists():
                        proposal = RosterRows.model_validate_json(cache.read_text())
                    else:
                        with Image.open(path.parent / "pages" / page.image) as image:
                            image.thumbnail(
                                (settings.reader_max_image_px, settings.reader_max_image_px)
                            )
                            buffer = io.BytesIO()
                            image.convert("RGB").save(buffer, "PNG")
                        raw = client.chat_json(
                            model.name,
                            prompt.render(),
                            RosterRows.model_json_schema(),
                            images=[buffer.getvalue()],
                        )
                        proposal = RosterRows.model_validate_json(raw)
                        cache.write_text(proposal.model_dump_json(), encoding="utf-8")
                    all_rows.extend(proposal.rows)
                item.rows = RosterRows(rows=all_rows).model_dump()["rows"]
            item.status, item.error = "ready", None
            db.commit()
        except (ValueError, UnicodeError, zipfile.BadZipFile, ValidationError):
            db.rollback()
            item = db.get_one(RosterImport, int(payload["roster_id"]))
            item.status, item.error = "manual", "roster.invalid"
            db.commit()


def confirm(db: Session, actor: User, ident: int, data: RosterRows) -> list[Student]:
    item = visible(db, actor, ident)
    if item.status == "confirmed":
        raise ConflictError("Register already confirmed")
    if item.status == "queued":
        raise ConflictError("Wait for extraction")
    nonempty = [row.code for row in data.rows if row.code]
    if len(nonempty) != len(set(nonempty)):
        raise ValidationFailedError("Duplicate code", code="roster.duplicate_code")
    result = []
    for index, row in enumerate(data.rows):
        if row.code:
            student = db.scalar(
                select(Student).where(
                    Student.school_id == actor.school_id, Student.code == row.code
                )
            )
            if not student:
                raise NotFoundError("Student code")
            get_visible_student(db, actor, student.id)
            student.class_id = item.class_id
            student.display_name = row.name or None
            student.is_active = True
        else:
            student = students.add_students(db, actor, item.class_id, [row.name])[0]
        student.register_position = index + 1
        result.append(student)
    item.rows = data.model_dump()["rows"]
    item.status = "confirmed"
    audit.record(
        db,
        actor=actor,
        entity_type="roster_import",
        entity_id=item.id,
        action="confirm_register",
        new_value={"student_ids": [s.id for s in result]},
    )
    return result


def eligible_student(exam: Exam, student: Student) -> bool:
    return student.school_id == exam.school_class.school_id and (
        student.id in (exam.roster_student_ids or [])
        or (student.class_id == exam.class_id and student.is_active)
    )
