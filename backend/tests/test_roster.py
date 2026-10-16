import io
from pathlib import Path
from typing import Any

from fastapi import FastAPI
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import Exam, Student
from app.services.roster import exam_students
from tests.conftest import World
from tests.factories import ApiUser, Factory, make_ready_exam, run_jobs, start_ai
from tests.test_review import _scan_exam


def upload(user: ApiUser, class_id: int, content: bytes, suffix: str) -> dict[str, Any]:
    response = user.post(f"/api/classes/{class_id}/imports?suffix={suffix}", content=content)
    assert response.status_code == 202, response.text
    run_jobs(user.app)
    return next(
        r
        for r in user.get(f"/api/classes/{class_id}/imports").json()
        if r["id"] == response.json()["id"]
    )


def test_csv_draft_confirm_reuses_identity_and_teacher_scope(
    app: FastAPI, db: Session, world: World, factory: Factory
) -> None:
    old = factory.student(world.class_a, "AAAAAA", "Old name")
    alice = ApiUser(app, "alice")
    draft = upload(
        alice, world.class_a.id, b"Name,Code\nCorrected name,AAAAAA\nNew student,\n", ".csv"
    )
    assert draft["status"] == "ready"
    assert len(list(db.scalars(select(Student)))) == 1
    response = alice.post(
        f"/api/roster-imports/{draft['id']}/confirm", json={"rows": draft["rows"]}
    )
    assert response.status_code == 200, response.text
    assert response.json()[0]["id"] == old.id
    db.expire_all()
    assert old.display_name == "Corrected name"
    assert len(list(db.scalars(select(Student)))) == 2
    assert (
        alice.post(
            f"/api/roster-imports/{draft['id']}/confirm", json={"rows": draft["rows"]}
        ).status_code
        == 409
    )
    bruno = ApiUser(app, "bruno")
    assert bruno.get(f"/api/classes/{world.class_a.id}/imports").status_code == 404
    assert bruno.get(f"/api/roster-imports/{draft['id']}/original").status_code == 404
    assert (
        alice.post(f"/api/classes/{world.class_b.id}/students", json={"names": ["No"]}).status_code
        == 404
    )
    assert (
        alice.post(f"/api/classes/{world.class_a.id}/students", json={"names": ["Yes"]}).status_code
        == 201
    )


def test_xlsx_and_pasted_names_preview_and_formula_rejection(app: FastAPI, world: World) -> None:
    alice = ApiUser(app, "alice")
    book = Workbook()
    sheet = book.active
    assert sheet
    sheet.append(["Nom", "Code"])
    sheet.append(["Élève Un", ""])
    sheet.append(["Élève Deux", ""])
    buffer = io.BytesIO()
    book.save(buffer)
    draft = upload(alice, world.class_a.id, buffer.getvalue(), ".xlsx")
    assert draft["rows"] == [{"name": "Élève Un", "code": ""}, {"name": "Élève Deux", "code": ""}]
    pasted = upload(alice, world.class_a.id, b"Doe, Jane\nSecond Person", ".txt")
    assert pasted["rows"][0]["name"] == "Doe, Jane"
    sheet["A2"] = '=HYPERLINK("https://outside.invalid","bad")'
    buffer = io.BytesIO()
    book.save(buffer)
    rejected = upload(alice, world.class_a.id, buffer.getvalue(), ".xlsx")
    assert rejected["status"] == "manual"
    assert rejected["error"] == "roster.invalid"


def test_register_order_in_answer_sheets_and_removed_student_history(
    app: FastAPI, db: Session, world: World, factory: Factory
) -> None:
    alice = ApiUser(app, "alice")
    first = factory.student(world.class_a, "AAAAAA")
    second = factory.student(world.class_a, "BBBBBB")
    assert (
        alice.put(
            f"/api/classes/{world.class_a.id}/register-order",
            json={"student_ids": [second.id, first.id]},
        ).status_code
        == 204
    )
    exam = make_ready_exam(alice, world.class_a.id)
    result = alice.post(
        f"/api/exams/{exam['id']}/answer-sheets",
        json={"pages_per_student": 1, "spares": 0, "request_key": "register-order-123456"},
    )
    assert result.status_code == 201, result.text
    db.expire_all()
    found = db.get_one(Exam, exam["id"])
    assert [s.id for s in exam_students(db, found)] == [second.id, first.id]
    assert alice.put(f"/api/students/{first.id}/active", json={"active": False}).status_code == 200
    db.expire_all()
    assert db.get_one(Student, first.id).is_active is False
    assert alice.get(f"/api/students/{first.id}/history").status_code == 200
    assert ApiUser(app, "bruno").get(f"/api/students/{first.id}/history").status_code == 404


def test_excel_preserves_numeric_marks_escapes_formulas_and_history_after_transfer(
    app: FastAPI, db: Session, world: World, factory: Factory, tmp_path: Path
) -> None:
    alice = ApiUser(app, "alice")
    exam_id = _scan_exam(alice, world, factory, ("AAAAAA", "BBBBBB"))
    start_ai(alice, exam_id)
    run_jobs(app)
    items = alice.get(f"/api/exams/{exam_id}/queue").json()["items"]
    for item in items:
        assert alice.post(f"/api/scripts/{item['script_id']}/accept-all").status_code == 200
        assert alice.post(f"/api/scripts/{item['script_id']}/confirm").status_code == 200
    student = db.scalar(select(Student).where(Student.code == "AAAAAA"))
    assert student
    student.display_name = "=2+2"
    db.commit()
    assert alice.put(f"/api/exams/{exam_id}/term", json={"term": "Term 1"}).status_code == 204
    exported = alice.get(f"/api/exams/{exam_id}/results.xlsx")
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content))
    sheet = book.active
    assert sheet
    assert sheet["B2"].value == "'=2+2"
    assert sheet["F2"].value == 2
    assert sheet["F2"].data_type == "n"
    assert sheet["D2"].value == "Term 1"
    student.class_id = world.class_b.id
    db.commit()
    history = alice.get(f"/api/students/{student.id}/history")
    assert history.status_code == 200
    assert history.json()[0]["term"] == "Term 1"
    assert history.json()[0]["total"] == 3
    # The old exam's roster survives the move; no private new-class exam is included.
    assert [r["exam_id"] for r in history.json()] == [exam_id]


def test_register_photo_uses_only_local_model_and_requires_confirmation(
    app: FastAPI, world: World, db: Session
) -> None:
    import json

    from PIL import Image

    from tests.fake_ollama import FakeOllama

    alice = ApiUser(app, "alice")
    buffer = io.BytesIO()
    Image.new("RGB", (500, 700), "white").save(buffer, "PNG")
    with FakeOllama() as local:
        local.default_reply = json.dumps({"rows": [{"name": "Read locally", "code": ""}]})
        app.state.settings.ai_backend = "ollama"
        app.state.settings.model_base_url = local.url
        draft = upload(alice, world.class_a.id, buffer.getvalue(), ".png")
        assert draft["status"] == "ready"
        assert draft["model_name"] == "qwen3-vl:4b-instruct"
        assert draft["prompt_version"].startswith("register-read-v1#")
        assert len(local.requests) == 1
    assert list(db.scalars(select(Student))) == []
