from pathlib import Path
from typing import Any

import pytest
from docx import Document
from fastapi import FastAPI
from sqlalchemy.orm import Session

from app.config import Settings
from app.db.models import Exam, MarkSchemePoint, Question
from app.schemas.mark_scheme import MarkSchemeIn
from app.services.ai_marking import exam_spec
from app.services.mark_scheme import validate_mark_scheme
from app.services.scheme_documents import docx_text, extract
from app.services.scoring import score_question
from tests.conftest import World
from tests.factories import ApiUser, complete_scheme


def question() -> Question:
    return Question(
        id=10,
        label="1(a)",
        max_half_marks=8,
        groups=[{"key": "facts", "max_points": 2, "max_marks": 2}],
        points=[
            MarkSchemePoint(id=1, half_marks=2, rules={"key": "M1", "kind": "M"}),
            MarkSchemePoint(
                id=2, half_marks=2, rules={"key": "A1", "kind": "A", "depends_on": ["M1"]}
            ),
            *[MarkSchemePoint(id=n, half_marks=2, rules={"group": "facts"}) for n in (3, 4, 5)],
        ],
    )


def test_scoring_applies_method_dependencies_and_any_n_caps() -> None:
    q = question()
    assert score_question(q, {1: 0, 2: 2, 3: 2, 4: 2, 5: 2}) == 4
    assert score_question(q, {1: 2, 2: 2, 3: 2, 4: 2, 5: 2}) == 8
    assert score_question(q, {}) is None
    assert score_question(q, {1: 0, 2: 0}) == 0


@pytest.mark.parametrize(
    "rules",
    [
        {"kind": "A"},
        {"key": "a", "depends_on": ["missing"]},
        {"key": "a", "depends_on": ["a"]},
        {"group": "missing"},
    ],
)
def test_invalid_dependencies_and_groups_block_saving(rules: dict[str, Any]) -> None:
    scheme = complete_scheme()
    scheme["questions"][0]["points"][0]["rules"] = rules
    report = validate_mark_scheme(MarkSchemeIn.model_validate(scheme))
    assert "scheme.rules_invalid" in [i.code for i in report.blocking]


def test_group_capacity_is_used_in_readiness() -> None:
    data = complete_scheme()
    q = data["questions"][0]
    q["max_marks"] = 2
    q["groups"] = [{"key": "g", "max_points": 1, "max_marks": 1}]
    for p in q["points"]:
        p["rules"] = {"group": "g"}
    report = validate_mark_scheme(MarkSchemeIn.model_validate(data))
    assert not report.is_ready
    assert "question.points_below_max" in [i.code for i in report.readiness]


def test_docx_extracts_paragraphs_and_table_cells(tmp_path: Path) -> None:
    path = tmp_path / "scheme.docx"
    doc = Document()
    doc.add_paragraph("1(a)(ii) Method M1")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Reject oxygen alone"
    table.cell(0, 1).text = "A1 depends on M1"
    doc.save(str(path))
    text = docx_text(path)
    assert "1(a)(ii) Method M1" in text
    assert "Reject oxygen alone | A1 depends on M1" in text


def test_import_never_confirms_and_original_is_scoped(
    app: FastAPI, world: World, settings: Settings, db: Session
) -> None:
    alice, bruno = ApiUser(app, "alice"), ApiUser(app, "bruno")
    exam = alice.post(
        "/api/exams",
        json={"class_id": world.class_a.id, "subject_code": "chemistry", "title": "Scheme"},
    ).json()
    url = f"/api/exams/{exam['id']}/scheme-documents?suffix=.txt&kind=scheme"
    response = alice.post(url, content=b"1(a) Oxygen: 1 mark")
    assert response.status_code == 202, response.text
    doc = response.json()
    assert alice.post(url, content=b"1(a) Oxygen: 1 mark").json()["id"] == doc["id"]
    original = f"/api/scheme-documents/{doc['id']}/original"
    assert alice.get(original).content == b"1(a) Oxygen: 1 mark"
    assert bruno.get(original).status_code == 404
    assert bruno.get(f"/api/exams/{exam['id']}/scheme-documents").status_code == 404
    assert alice.client.post(url, content=b"bad").status_code == 403
    extract(app.state.sessionmaker, settings, {"document_id": doc["id"]})
    docs = alice.get(f"/api/exams/{exam['id']}/scheme-documents").json()
    assert docs[0]["status"] == "manual"
    assert docs[0]["proposal"] is None
    assert alice.get(f"/api/exams/{exam['id']}").json()["status"] == "draft"
    assert alice.post("/api/scheme-library", json={"exam_id": exam["id"]}).status_code == 409
    scheme = complete_scheme()
    scheme["questions"][0]["points"][0]["rules"] = {"teacher_only": True}
    assert alice.put(f"/api/exams/{exam['id']}/mark-scheme", json=scheme).status_code == 200
    assert alice.put(f"/api/exams/{exam['id']}/status", json={"status": "ready"}).status_code == 200
    entry = alice.post("/api/scheme-library", json={"exam_id": exam["id"]})
    assert entry.status_code == 201
    shared = bruno.get(f"/api/scheme-library/{entry.json()['id']}")
    assert shared.json()["questions"][0]["points"][0]["rules"]["teacher_only"] is True
    stored = db.get_one(Exam, exam["id"])
    spec = exam_spec(stored, world.school.curriculum)
    assert len(spec.questions[0].points) == 1
    assert all(not p.rules.teacher_only for q in spec.questions for p in q.points)


def test_local_extraction_proposal_never_changes_saved_scheme(
    app: FastAPI, world: World, settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.ai.types import ModelInfo
    from app.services import scheme_documents

    class Client:
        def chat_json(
            self, model: str, prompt: str, schema: dict[str, Any], images: list[bytes] | None = None
        ) -> str:
            assert "DOCUMENT:" in prompt
            assert "Hydrogen" in prompt
            return MarkSchemeIn.model_validate(complete_scheme()).model_dump_json()

    monkeypatch.setattr(scheme_documents, "ollama_client", lambda _: Client())
    monkeypatch.setattr(
        scheme_documents, "resolve_model", lambda *_: ModelInfo("local", "digest", "ollama")
    )
    alice = ApiUser(app, "alice")
    exam = alice.post(
        "/api/exams",
        json={"class_id": world.class_a.id, "subject_code": "chemistry", "title": "Draft"},
    ).json()
    doc = alice.post(
        f"/api/exams/{exam['id']}/scheme-documents?suffix=.txt&kind=scheme",
        content=b"Hydrogen 1 mark",
    ).json()
    settings.ai_backend = "ollama"
    extract(app.state.sessionmaker, settings, {"document_id": doc["id"]})
    stored = alice.get(f"/api/exams/{exam['id']}/scheme-documents").json()[0]
    assert stored["status"] == "ready"
    assert stored["proposal"]["questions"][0]["label"] == "1(a)"
    assert stored["model_version"] == "digest"
    assert stored["prompt_version"].startswith("scheme-extract-v1#")
    actual = alice.get(f"/api/exams/{exam['id']}").json()
    assert actual["questions"] == []
    assert actual["status"] == "draft"
